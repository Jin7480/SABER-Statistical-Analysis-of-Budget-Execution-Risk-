# -*- coding: utf-8 -*-
"""
saber_model.py — 세이버-K(SABER-K): 계층적 정규화 기반 예산 집행률 예측 모형
Statistical Analysis of Budget Execution Risk

핵심 구성 (경험적 베이즈 근사 구현):
  1) 교차 계층 앵커 M: 전체 평균 + (동일 사업명 효과) + (소속 기관 효과)
     — 이력이 없는 신규 사업도 상위 계층 정보로 예측 (콜드 스타트 해소)
  2) 최신성 가중 개별 신호 m_ew: 지수 감쇠(λ) 가중 평균 — 직전 연도 정보를 상태로 통합
  3) 정보량 비례 축소: 이력 길이 구간별 축소계수 β를 '검증 연도'에서 회귀로 추정
     — 이력이 짧을수록 앵커(집단 정보) 비중이 커지는 부분 풀링(partial pooling)
  예측: pred = M + β_last(구간)·(last − M) + β_ew(구간)·(m_ew − M),  신규(n=0)는 pred = M
"""
import numpy as np
import pandas as pd

HIST_BINS = [(1, 1), (2, 2), (3, 4), (5, 6), (7, 99)]  # 이력 길이 구간


class SaberModel:
    def __init__(self, lam: float = 0.5, cP: float = 20, cG: float = 200):
        self.lam, self.cP, self.cG = lam, cP, cG
        self.betas = None  # 구간별 (β_last, β_ew)

    # ---------- 내부: 학습 데이터 통계 ----------
    def _stats(self, past: pd.DataFrame, train_end: int):
        past = past.sort_values('연도').copy()
        past['wt'] = self.lam ** (train_end - past['연도'])
        past['wy'] = past['wt'] * past['집행률']
        g = past.groupby('uid')
        unit = pd.DataFrame({
            'n': g['집행률'].size(),
            'm_ew': g['wy'].sum() / g['wt'].sum(),
            'last': g['집행률'].last()})
        aP = past.groupby('사업명')['집행률'].agg(['mean', 'size']); aP.columns = ['mP', 'sP']
        aG = past.groupby(['지역', '자치단체'])['집행률'].agg(['mean', 'size']); aG.columns = ['mG', 'sG']
        return past['집행률'].mean(), unit, aP, aG

    def _features(self, rows: pd.DataFrame, GLOBAL, unit, aP, aG) -> pd.DataFrame:
        t = rows.merge(unit, left_on='uid', right_index=True, how='left')
        t['n'] = t['n'].fillna(0).astype(int)
        t = t.merge(aP, left_on='사업명', right_index=True, how='left')
        t = t.merge(aG, left_on=['지역', '자치단체'], right_index=True, how='left')
        bP = t['sP'].fillna(0) / (t['sP'].fillna(0) + self.cP)
        bG = t['sG'].fillna(0) / (t['sG'].fillna(0) + self.cG)
        t['M'] = GLOBAL + bP * (t['mP'].fillna(GLOBAL) - GLOBAL) \
                        + bG * (t['mG'].fillna(GLOBAL) - GLOBAL)
        t['nb'] = t['n'].map(self._nbin)
        return t

    @staticmethod
    def _nbin(n):
        for i, (a, b) in enumerate(HIST_BINS):
            if a <= n <= b:
                return i
        return len(HIST_BINS) - 1

    # ---------- 학습: 검증 연도에서 구간별 축소계수 회귀 추정 ----------
    def fit(self, df: pd.DataFrame, train_end: int, val_year: int):
        """train_end 이전 자료를 학습, val_year 자료로 축소계수 β를 추정.
        시험 연도 데이터는 일절 접촉하지 않음 (시간 분할 검증 위생)."""
        GLOBAL, unit, aP, aG = self._stats(df[df.연도 <= train_end], train_end)
        val = self._features(df[df.연도 == val_year].copy(), GLOBAL, unit, aP, aG)
        val = val[val.n >= 1].dropna(subset=['last', 'm_ew'])
        self.betas = {}
        for i in range(len(HIST_BINS)):
            v = val[val.nb == i]
            X = np.column_stack([v['last'] - v['M'], v['m_ew'] - v['M']])
            y = (v['집행률'] - v['M']).values
            self.betas[i], *_ = np.linalg.lstsq(X, y, rcond=None)
        return self

    # ---------- 예측 ----------
    def predict(self, df: pd.DataFrame, train_end: int, test_year: int) -> pd.DataFrame:
        assert self.betas is not None, 'fit() 먼저 호출'
        GLOBAL, unit, aP, aG = self._stats(df[df.연도 <= train_end], train_end)
        t = self._features(df[df.연도 == test_year].copy(), GLOBAL, unit, aP, aG)
        b0 = np.array([self.betas[i][0] for i in range(len(HIST_BINS))])
        b1 = np.array([self.betas[i][1] for i in range(len(HIST_BINS))])
        t['pred_saber'] = np.where(
            t.n >= 1,
            t['M'] + b0[t['nb']] * (t['last'] - t['M']).fillna(0)
                   + b1[t['nb']] * (t['m_ew'] - t['M']).fillna(0),
            t['M'])  # 신규 사업: 교차 계층 앵커로 예측 (기존 방식은 예측 불가)
        return t

    # ---------- 예측구간: 이력 구간별 잔차 산포로 근사 ----------
    def fit_intervals(self, df: pd.DataFrame, train_end: int, val_year: int,
                      coverage: float = 0.8):
        """검증 연도의 구간별 예측 잔차 표준편차로 예측구간 폭을 추정한다.
        경험적 베이즈 근사 모델에는 사후 분포가 없으므로, 잔차 기반으로
        '이력이 짧을수록 넓은' 예측구간을 산출한다(대시보드의 80% 구간 근거).

        coverage: 목표 예측구간 커버리지(예: 0.8 → 80% 구간).
        학습된 self.interval_half[구간] = 반폭(half-width)을 저장한다.
        """
        from scipy.stats import norm
        assert self.betas is not None, 'fit() 먼저 호출'
        z = norm.ppf(0.5 + coverage / 2.0)  # 예: 0.8 → 1.2816
        pred = self.predict(df, train_end - 1, val_year)  # 검증 연도 예측
        pred = pred[pred.n >= 1].copy()
        pred['resid'] = pred['집행률'] - pred['pred_saber']
        self.interval_half = {}
        for i in range(len(HIST_BINS)):
            r = pred.loc[pred.nb == i, 'resid'].dropna()
            sd = r.std() if len(r) >= 30 else pred['resid'].std()
            self.interval_half[i] = float(z * sd)
        self.interval_coverage = coverage
        return self

    def predict_with_interval(self, df: pd.DataFrame, train_end: int,
                              test_year: int) -> pd.DataFrame:
        """예측값에 하한/상한(예측구간)을 붙여 반환한다. fit_intervals() 선행 필요."""
        assert getattr(self, 'interval_half', None) is not None, \
            'fit_intervals() 먼저 호출'
        t = self.predict(df, train_end, test_year)
        half = t['nb'].map(self.interval_half).astype(float)
        t['pred_low'] = (t['pred_saber'] - half).clip(0, 1)
        t['pred_high'] = (t['pred_saber'] + half).clip(0, 1)
        t['interval_half'] = half
        return t

    # ---------- 위험도 산정: 예상 불용액·등급·우선순위 ----------
    @staticmethod
    def risk_table(pred: pd.DataFrame,
                   budget_col: str = '예산현액',
                   exclude_kw: tuple = ('보전지출', '내부거래', '예치', '기금', '예비비'),
                   min_budget: float = 3e8,
                   risk_threshold: float = 0.75) -> pd.DataFrame:
        """예측 결과로부터 재배정 심의용 위험도 표를 만든다.

        - 예상 불용액 = 예산현액 × (1 − 예측 집행률)
        - 위험 등급: 예측 집행률 < 0.4 심각 / < 0.6 높음 / < risk_threshold 주의
        - 회계적 이전지출(보전지출 등)과 소액 사업은 제외
        - 예상 불용액 내림차순 정렬(재배정 실익 큰 순)

        반환 컬럼: [사업명, 예산현액, 예측집행률, 예상불용액, 위험등급, uid ...]
        """
        t = pred.copy()
        if exclude_kw:
            mask = ~t['사업명'].astype(str).str.contains('|'.join(exclude_kw), na=False)
            t = t[mask]
        t = t[t[budget_col] >= min_budget]
        t = t[t['pred_saber'] < risk_threshold].copy()
        t['예상불용액'] = t[budget_col] * (1 - t['pred_saber'])

        def _grade(p):
            if p < 0.4:
                return '심각'
            if p < 0.6:
                return '높음'
            return '주의'
        t['위험등급'] = t['pred_saber'].map(_grade)
        t = t.sort_values('예상불용액', ascending=False)
        keep = ['지역', '자치단체', '사업명', budget_col, 'pred_saber',
                '예상불용액', '위험등급', 'n']
        keep = [c for c in keep if c in t.columns]
        out = t[keep].rename(columns={'pred_saber': '예측집행률', 'n': '보유이력'})
        return out.reset_index(drop=True)

    @staticmethod
    def field_check_priority(pred: pd.DataFrame, min_history: int = 2) -> pd.DataFrame:
        """예측 불확실성이 큰(이력이 짧은) 사업을 '현장 확인 우선' 대상으로 표시.
        interval_half가 있으면 그 값 기준, 없으면 이력 길이 기준으로 판정한다."""
        t = pred.copy()
        if 'interval_half' in t.columns:
            thr = t['interval_half'].quantile(0.75)
            t['현장확인우선'] = t['interval_half'] >= thr
        else:
            t['현장확인우선'] = t['n'] <= min_history
        return t

    # ---------- 진단: 집행이 유의하게 지연되는 부대/기관 자동 탐지 ----------
    @staticmethod
    def diagnose_laggards(df: pd.DataFrame, train_end: int,
                          group_col: str = '자치단체',
                          min_n: int = 30, z_thresh: float = 2.0) -> pd.DataFrame:
        """전체 평균 대비 집행률이 유의하게 낮은 기관을 무선효과 관점에서 탐지한다.
        기획서의 '제도·절차 병목 진단(행정 컨설팅)' 부수 기능의 실제 구현.

        각 기관 평균과 전체 평균의 차이를 표준오차로 나눈 z점수로 정렬하고,
        z ≤ −z_thresh(기본 −2.0)인 기관을 '유의 지연'으로 표시한다.
        min_n: 표본이 이보다 적은 기관은 제외(우연 변동 방지).
        """
        past = df[df['연도'] <= train_end]
        overall = past['집행률'].mean()
        g = past.groupby(group_col)['집행률']
        tab = pd.DataFrame({'평균집행률': g.mean(), 'n': g.size(), '표준편차': g.std()})
        tab = tab[tab['n'] >= min_n].copy()
        se = tab['표준편차'] / np.sqrt(tab['n'])
        tab['전체대비차'] = tab['평균집행률'] - overall
        tab['z'] = tab['전체대비차'] / se.replace(0, np.nan)
        tab['유의지연'] = tab['z'] <= -z_thresh
        tab = tab.sort_values('z')
        return tab.reset_index()

    # ---------- 모델 저장/불러오기 ----------
    def save(self, path: str):
        """학습된 하이퍼파라미터·축소계수·예측구간을 JSON으로 저장한다."""
        import json
        state = {
            'lam': self.lam, 'cP': self.cP, 'cG': self.cG,
            'betas': {str(k): list(map(float, v)) for k, v in (self.betas or {}).items()},
            'interval_half': getattr(self, 'interval_half', None),
            'interval_coverage': getattr(self, 'interval_coverage', None),
            'hist_bins': HIST_BINS,
        }
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(state, f, ensure_ascii=False, indent=2)
        return path

    @classmethod
    def load(cls, path: str) -> 'SaberModel':
        """save()로 저장한 모델 상태를 복원한다."""
        import json
        with open(path, encoding='utf-8') as f:
            state = json.load(f)
        m = cls(lam=state['lam'], cP=state['cP'], cG=state['cG'])
        m.betas = {int(k): np.array(v) for k, v in state['betas'].items()}
        if state.get('interval_half') is not None:
            m.interval_half = {int(k): v for k, v in state['interval_half'].items()}
            m.interval_coverage = state.get('interval_coverage')
        return m


# ============================================================
# 캐시 생성 유틸리티 (독립 실행용, 자동 호출되지 않음)
# ============================================================
def build_cache(data_dir: str = 'data',
                cache_path: str = 'data/panel_cache.pkl',
                min_budget: int = 1_000_000) -> 'pd.DataFrame':
    """원본 엑셀(연도별 세부사업별 세출현황)을 전처리해 패널 캐시를 생성·저장한다.

    이 함수는 필요할 때 사용자가 직접 호출하는 유틸리티이며, 모듈 임포트 시
    자동으로 실행되지 않는다. evaluate.py의 load_panel()과 동일한 전처리
    (열 선택 → 정제 → 사업 단위 집계 → 집행률 계산)를 수행한다.

    처리 순서:
      1) data_dir 안의 dataset_*.zip 이 있으면 압축 해제
      2) *_세부사업별세출현황.xlsx 를 연도별로 순차 파싱 (메모리 절약: read-only 스트리밍)
      3) 연도별로 정제·집계 후 합쳐 집행률·uid 계산
      4) cache_path 로 저장

    매개변수:
      data_dir    : 엑셀(및 zip)이 있는 폴더
      cache_path  : 저장할 캐시 경로(.pkl)
      min_budget  : 이 금액 미만의 미세 예산 항목은 제외 (기본 100만 원)

    반환: 전처리된 pandas.DataFrame (저장과 동일한 내용)

    사용 예:
      python -c "from saber_model import build_cache; build_cache()"
    """
    import os, glob, zipfile
    from openpyxl import load_workbook

    # 1) zip 자동 해제
    for z in glob.glob(os.path.join(data_dir, 'dataset_*.zip')):
        with zipfile.ZipFile(z) as f:
            f.extractall(data_dir)

    # 2~3) 연도별 파싱 → 정제 → 집계
    parts = []
    xlsx_list = sorted(glob.glob(os.path.join(data_dir, '*_세부사업별세출현황.xlsx')))
    if not xlsx_list:
        raise FileNotFoundError(
            f"{data_dir} 에서 '*_세부사업별세출현황.xlsx' 를 찾을 수 없습니다. "
            f"엑셀 또는 dataset_*.zip 을 먼저 넣어주세요.")
    for path in xlsx_list:
        year = int(os.path.basename(path)[:4])
        print(f'로드 중: {year} ...', flush=True)
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        recs = []
        for r in ws.iter_rows(min_row=3, values_only=True):
            if r[0] is None or r[3] is None:
                continue
            recs.append((r[0], r[1], r[2], r[3], r[4], r[9], r[11], r[12]))
        wb.close()
        d = pd.DataFrame(recs, columns=['지역', '자치단체', '회계', '사업명',
                                        '예산현액', '지출액', '분야', '부문'])
        del recs
        d['예산현액'] = pd.to_numeric(d['예산현액'], errors='coerce')
        d['지출액'] = pd.to_numeric(d['지출액'], errors='coerce')
        d = d.dropna(subset=['예산현액', '지출액'])
        d = d[d['예산현액'] >= min_budget]
        d = d.groupby(['지역', '자치단체', '회계', '사업명'], as_index=False).agg(
            {'예산현액': 'sum', '지출액': 'sum', '분야': 'first', '부문': 'first'})
        d['연도'] = year
        parts.append(d)
        del d

    # 4) 합치기 → 집행률·uid → 저장
    df = pd.concat(parts, ignore_index=True)
    del parts
    df['집행률'] = (df['지출액'] / df['예산현액']).clip(0, 1)
    df['uid'] = (df['지역'] + '|' + df['자치단체'] + '|'
                 + df['회계'].astype(str) + '|' + df['사업명'])
    os.makedirs(os.path.dirname(cache_path) or '.', exist_ok=True)
    df.to_pickle(cache_path)
    print(f'캐시 저장 완료: {cache_path}  ({len(df):,} 행)')
    return df




# ============================================================
# [향후 확장] 완전 계층 베이지안 모형 (로드맵 2단계 구현 예정)
# ------------------------------------------------------------
# 아래 두 함수는 향후 구축할 "완전한 확률적 계층 베이지안 회귀"의 참조 구현이다.
# 현재 검증에 사용한 SaberModel(위)은 이 베이지안 모형의 사후 평균을 닫힌형으로 근사한
# 경험적 베이즈(empirical Bayes) 버전이며, 대규모 데이터에서 GPU 없이 수 분 내 학습된다.
# 완전 베이지안 버전은 무선효과의 사후 분포와 예측 신뢰구간을 직접 표본화하지만
# 계산 비용이 크므로, 파일럿 검증이 아닌 실배치(로드맵 2단계)에서 사용한다.
#
# 두 모형의 관계:
#   - 반응변수: 집행률은 [0,1] 유계 → 베타 회귀,  불용액(원)은 0 과잉 + 양의 연속 → 영과잉 감마
#   - 계층: 사업명·기관·분야를 무선효과(random effect)로 두어 부분 풀링 수행
#   - SaberModel은 위 구조의 사후 평균을 축소추정으로 근사한 것이고,
#     아래 함수는 동일 계층 구조를 확률적으로 완전 추정한다(신뢰구간 산출 가능).
#
# 이 함수들은 모듈 임포트 시 실행되지 않으며, PyMC 등 별도 라이브러리가 설치된
# 환경에서 명시적으로 호출할 때만 동작한다. GitHub의 결과 재현(evaluate.py)에는
# 사용되지 않는다.
# ============================================================
def fit_full_bayesian_beta(df, train_end, target='집행률',
                           group_cols=('사업명', '자치단체', '분야'),
                           draws=1000, tune=1000, chains=2):
    """집행률(0~1 유계)에 대한 완전 계층 베이지안 '베타 회귀' 참조 구현.

    각 계층(사업명·기관·분야)에 무선효과를 두고, 집행률의 평균 mu를 로짓 링크로
    모델링한다. 사후 표본에서 예측 평균과 신뢰구간(예: 95% HDI)을 직접 얻는다.

    ※ 향후 확장용. PyMC가 설치된 환경에서만 동작하며, 현재 검증 파이프라인에서는
      호출되지 않는다. 계산 비용이 커 GPU 또는 충분한 CPU 시간이 필요하다.

    반환: pymc.InferenceData (사후 표본)
    """
    try:
        import pymc as pm
        import numpy as np
    except ImportError:
        raise ImportError(
            "fit_full_bayesian_beta 는 PyMC가 필요합니다: pip install pymc\n"
            "(이 함수는 향후 확장용이며, 결과 재현에는 SaberModel을 사용하세요.)")

    past = df[df['연도'] <= train_end].copy()
    y = past[target].clip(1e-4, 1 - 1e-4).values  # 베타 분포 지지역 안으로
    # 계층 인덱스 생성
    idx, levels = {}, {}
    for c in group_cols:
        codes, uniq = pd.factorize(past[c].astype(str))
        idx[c] = codes
        levels[c] = len(uniq)

    with pm.Model() as model:
        # 전역 절편
        intercept = pm.Normal('intercept', 0.0, 1.5)
        # 계층별 무선효과 (비중심화 파라미터화)
        re_terms = []
        for c in group_cols:
            sigma_c = pm.HalfNormal(f'sigma_{c}', 1.0)
            z_c = pm.Normal(f'z_{c}', 0.0, 1.0, shape=levels[c])
            re_terms.append((sigma_c * z_c)[idx[c]])
        eta = intercept + sum(re_terms)
        mu = pm.Deterministic('mu', pm.math.sigmoid(eta))
        phi = pm.HalfNormal('phi', 5.0)  # 베타 정밀도
        alpha = mu * phi
        beta = (1.0 - mu) * phi
        pm.Beta('obs', alpha=alpha, beta=beta, observed=y)
        idata = pm.sample(draws=draws, tune=tune, chains=chains,
                          target_accept=0.9, progressbar=True)
    return idata


def fit_full_bayesian_zig(df, train_end, target_amount='불용액',
                          group_cols=('사업명', '자치단체', '분야'),
                          draws=1000, tune=1000, chains=2):
    """불용액(0 과잉 + 양의 연속)에 대한 완전 계층 베이지안 '영과잉 감마' 참조 구현.

    불용 발생 여부(베르누이)와 발생 시 규모(감마)를 동시에 모델링하는 hurdle 구조.
    무선효과로 계층 정보를 공유하며, 사후 표본에서 예상 불용액과 신뢰구간을 얻는다.

    ※ 향후 확장용. PyMC 필요. 현재 검증 파이프라인에서는 호출되지 않는다.

    반환: pymc.InferenceData (사후 표본)
    """
    try:
        import pymc as pm
        import numpy as np
    except ImportError:
        raise ImportError(
            "fit_full_bayesian_zig 는 PyMC가 필요합니다: pip install pymc\n"
            "(이 함수는 향후 확장용이며, 결과 재현에는 SaberModel을 사용하세요.)")

    past = df[df['연도'] <= train_end].copy()
    if target_amount not in past.columns:
        past = past.assign(불용액=(past['예산현액'] * (1 - past['집행률'])).clip(lower=0))
    amt = past[target_amount].values
    is_unspent = (amt > 0).astype(int)
    pos_amt = np.where(amt > 0, amt, 1.0)  # 감마 지지역(>0) 보정

    idx, levels = {}, {}
    for c in group_cols:
        codes, uniq = pd.factorize(past[c].astype(str))
        idx[c] = codes
        levels[c] = len(uniq)

    with pm.Model() as model:
        # (1) 불용 발생 여부: 로지스틱 + 계층 무선효과
        a0 = pm.Normal('a0', 0.0, 1.5)
        p_terms = []
        for c in group_cols:
            s = pm.HalfNormal(f'p_sigma_{c}', 1.0)
            z = pm.Normal(f'p_z_{c}', 0.0, 1.0, shape=levels[c])
            p_terms.append((s * z)[idx[c]])
        p_unspent = pm.Deterministic('p_unspent', pm.math.sigmoid(a0 + sum(p_terms)))
        pm.Bernoulli('occur', p=p_unspent, observed=is_unspent)

        # (2) 발생 시 규모: 감마 (로그 링크 + 계층 무선효과)
        b0 = pm.Normal('b0', np.log(pos_amt.mean() + 1.0), 2.0)
        m_terms = []
        for c in group_cols:
            s = pm.HalfNormal(f'm_sigma_{c}', 1.0)
            z = pm.Normal(f'm_z_{c}', 0.0, 1.0, shape=levels[c])
            m_terms.append((s * z)[idx[c]])
        mu_amt = pm.Deterministic('mu_amt', pm.math.exp(b0 + sum(m_terms)))
        shape_g = pm.HalfNormal('shape_g', 2.0)
        mask = is_unspent.astype(bool)
        pm.Gamma('amount', alpha=shape_g, beta=shape_g / mu_amt[mask],
                 observed=pos_amt[mask])

        idata = pm.sample(draws=draws, tune=tune, chains=chains,
                          target_accept=0.9, progressbar=True)
    return idata


# ============================================================
# 위험도 대시보드 렌더링 (HTML 생성)
# ------------------------------------------------------------
# risk_table()의 출력을 받아 내부망 웹 대시보드(HTML)를 생성한다.
# 외부 라이브러리 없이 표준 라이브러리만 사용하며, 결과 HTML은 브라우저로 바로 열린다.
# 모듈 임포트 시 실행되지 않으며, 필요할 때 명시적으로 호출한다.
#   예)  df ... ; m = SaberModel().fit(...); pred = m.predict(...)
#        rt = SaberModel.risk_table(pred_seocho)
#        render_dashboard(rt, org='서울특별시 서초구', out_path='dashboard.html')
# evaluate.py(결과 재현)에서는 호출되지 않는다.
# ============================================================
def render_dashboard(risk_df, org: str = '', as_of: str = '',
                     out_path: str = 'dashboard.html', top_n: int = 30) -> str:
    """위험도 표(risk_table 출력)를 HTML 대시보드로 렌더링해 파일로 저장한다.

    risk_df : SaberModel.risk_table()이 반환한 DataFrame
              (사업명·예산현액·예측집행률·예상불용액·위험등급·보유이력 등)
    org     : 화면 우상단에 표시할 기관명(예: '서울특별시 서초구')
    as_of   : 기준 시점 문구(예: '2025년 2분기 말 기준')
    out_path: 저장할 HTML 경로
    top_n   : 표에 표시할 상위 사업 수

    반환: out_path
    """
    import html as _html

    df = risk_df.head(top_n).copy()
    budget_col = '예산현액' if '예산현액' in df.columns else df.columns[3]

    grade_class = {'심각': 'g-red', '높음': 'g-org', '주의': 'g-yel'}
    rows_html = []
    for rank, (_, r) in enumerate(df.iterrows(), 1):
        name = _html.escape(str(r.get('사업명', '')))
        gu = _html.escape(str(r.get('자치단체', '')))
        budget_eok = r[budget_col] / 1e8
        rate = float(r['예측집행률']) * 100
        unspent_eok = r['예상불용액'] / 1e8
        grade = str(r['위험등급'])
        gcls = grade_class.get(grade, 'g-yel')
        hist = int(r['보유이력']) if '보유이력' in df.columns and not pd.isna(r['보유이력']) else 0
        act = '현장 확인 우선 (이력 짧음·불확실성 큼)' if hist <= 2 else \
              ('계약·시공 일정 점검' if any(k in name for k in ('건립', '신축', '건설', '조성'))
               else '집행 계획 재점검')
        low = high = None
        if 'pred_low' in r and 'pred_high' in r and not pd.isna(r.get('pred_low')):
            low, high = float(r['pred_low']) * 100, float(r['pred_high']) * 100
        ci = f' <span class="ci">({low:.0f}~{high:.0f}%)</span>' if low is not None else ''
        rows_html.append(f'''
      <tr>
        <td class="num">{rank}</td>
        <td class="name">{name}<span class="sub">{gu}</span></td>
        <td>{budget_eok:,.0f}억</td>
        <td><span class="rate">{rate:.0f}%</span>{ci}
            <span class="bar"><i style="width:{rate:.0f}%"></i></span></td>
        <td class="amt">{unspent_eok:,.0f}억</td>
        <td><span class="badge {gcls}">{grade}</span></td>
        <td class="act">{_html.escape(act)}</td>
      </tr>''')

    total_unspent = df['예상불용액'].sum() / 1e8
    n_risk = len(risk_df)
    org_html = _html.escape(org)
    as_of_html = _html.escape(as_of) if as_of else ''

    doc = f'''<!DOCTYPE html>
<html lang="ko"><head><meta charset="utf-8">
<title>세이버-K | 예산 불용 조기경보</title>
<style>
  *{{margin:0;padding:0;box-sizing:border-box;}}
  body{{font-family:'Malgun Gothic','Noto Sans CJK KR',sans-serif;background:#eef1f6;padding:24px;color:#1d2b3d;}}
  .app{{max-width:1280px;margin:0 auto;background:#fff;border-radius:12px;
        box-shadow:0 6px 24px rgba(30,45,70,.12);overflow:hidden;}}
  .header{{background:#22354f;padding:18px 26px;display:flex;justify-content:space-between;align-items:center;}}
  .logo{{font-size:20px;font-weight:700;color:#fff;}}
  .logo .sub{{font-size:13px;font-weight:400;color:#9fb2cc;margin-left:12px;}}
  .org{{text-align:right;color:#fff;font-size:14px;font-weight:700;}}
  .org .meta{{display:block;font-size:11.5px;font-weight:400;color:#9fb2cc;margin-top:3px;}}
  .summary{{padding:14px 26px;color:#54657e;font-size:13px;border-bottom:1px solid #eef1f6;}}
  .summary b{{color:#c0392b;font-size:15px;}}
  table{{width:100%;border-collapse:collapse;}}
  th{{background:#eef2f8;color:#33445c;font-size:12px;text-align:left;padding:10px 12px;}}
  td{{padding:11px 12px;font-size:13px;border-bottom:1px solid #eef1f6;vertical-align:middle;}}
  .num{{color:#54657e;}} .name{{font-weight:600;}}
  .name .sub{{display:block;font-weight:400;color:#8a97a8;font-size:11px;}}
  .rate{{color:#4c72b0;font-weight:700;}} .ci{{color:#8a97a8;font-size:11px;}}
  .bar{{display:inline-block;width:70px;height:8px;background:#e2e7ee;border-radius:5px;margin-left:6px;vertical-align:middle;}}
  .bar i{{display:block;height:8px;background:#4c72b0;border-radius:5px;}}
  .amt{{color:#c0392b;font-weight:700;}}
  .badge{{display:inline-block;color:#fff;font-size:11px;font-weight:700;padding:3px 9px;border-radius:9px;}}
  .g-red{{background:#c0392b;}} .g-org{{background:#e67e22;}} .g-yel{{background:#b7950b;}}
  .act{{color:#54657e;font-size:11.5px;}}
  .foot{{padding:12px 26px 18px;color:#8a97a8;font-size:11px;}}
</style></head>
<body><div class="app">
  <div class="header">
    <div class="logo">세이버-K<span class="sub">예산 불용 조기경보</span></div>
    <div class="org">{org_html}<span class="meta">{as_of_html} · 분기 자동 갱신</span></div>
  </div>
  <div class="summary">불용 위험 사업 <b>{n_risk}건</b> · 3분기 재배정 전환 가능 재원(상위 {min(top_n, n_risk)}건 합계) 약 <b>{total_unspent:,.0f}억 원</b></div>
  <table>
    <tr><th>순위</th><th>세부사업명</th><th>예산현액</th><th>예측 집행률</th>
        <th>예상 불용액</th><th>위험</th><th>권장 조치</th></tr>{''.join(rows_html)}
  </table>
  <div class="foot">※ 예측치는 세이버-K 모형 산출값. 최종 판단은 담당자·심의기구가 수행. 회계적 이전지출은 제외됨.</div>
</div></body></html>'''

    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(doc)
    return out_path




if __name__ == '__main__':
    # 이 파일을 직접 실행하면 캐시를 생성한다:  python saber_model.py
    build_cache()
