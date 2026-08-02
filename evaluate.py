# -*- coding: utf-8 -*-
"""
evaluate.py — 데이터 로드 → 백테스트 → 성능 표·그림 생성
사용법:
  1) data/ 폴더에 dataset_*.zip 압축 해제 (연도별 xlsx 10개)
  2) python evaluate.py
출력: results/ 폴더에 성능 표(csv)와 그림(png)

백테스트 설계: 2016~2023 학습 → 2024 검증(축소계수 추정) → 2025 시험
"""
import os, glob, zipfile
import numpy as np
import pandas as pd

from baseline_models import predict_baselines
from saber_model import SaberModel

DATA_DIR, RESULT_DIR, CACHE = 'data', 'results', 'data/panel_cache.pkl'
TRAIN_END, VAL_YEAR, TEST_YEAR = 2023, 2024, 2025
BINS = [(0, 0), (1, 2), (3, 4), (5, 6), (7, 9)]
LABELS = ['신규(0년)', '1~2년', '3~4년', '5~6년', '7~9년']

# 원본 데이터(엑셀 zip 5개 + 전처리 캐시)가 담긴 Google Drive 공개 폴더 ID.
# GitHub 용량 제한 때문에 데이터는 Drive에 두고, 실행 시 자동으로 내려받는다.
GDRIVE_FOLDER_ID = '1455r3kIBEtm6gb0WTuAqWj4W6CmHADhR'


def _ensure_data():
    """data/ 폴더에 캐시나 원본이 없으면 Google Drive 공개 폴더에서 자동 다운로드한다."""
    os.makedirs(DATA_DIR, exist_ok=True)
    has_cache = os.path.exists(CACHE) or os.path.exists(CACHE + '.gz')
    has_xlsx = bool(glob.glob(f'{DATA_DIR}/*_세부사업별세출현황.xlsx'))
    has_zip = bool(glob.glob(f'{DATA_DIR}/dataset_*.zip'))
    if has_cache or has_xlsx or has_zip:
        return  # 이미 데이터가 있으면 다운로드하지 않음
    try:
        import gdown
    except ImportError:
        raise SystemExit(
            "데이터가 없고 gdown이 설치되어 있지 않습니다.\n"
            "  pip install gdown  을 실행하거나,\n"
            "  Google Drive 폴더에서 파일을 직접 받아 data/ 에 넣어주세요.")
    print('data/ 가 비어 있어 Google Drive에서 데이터를 내려받습니다 ...', flush=True)
    gdown.download_folder(id=GDRIVE_FOLDER_ID, output=DATA_DIR,
                          quiet=False, use_cookies=False)


# ---------- 데이터 준비 ----------
def load_panel() -> pd.DataFrame:
    """연도별 xlsx를 순차 파싱→즉시 축약→디스크에 임시 저장한 뒤 마지막에 합쳐,
    저사양(메모리 4GB급) 환경에서도 안전하게 전체 패널을 구성한다.
    캐시(panel_cache.pkl)가 있으면 파싱을 건너뛴다."""
    _ensure_data()  # data/ 가 비어 있으면 Google Drive에서 자동 다운로드
    if os.path.exists(CACHE):
        return pd.read_pickle(CACHE)
    # 배포용 압축 캐시(panel_cache.pkl.gz)가 있으면 그것을 사용 (엑셀 파싱 불필요)
    gz = CACHE + '.gz'
    if os.path.exists(gz):
        import gzip, pickle
        with gzip.open(gz, 'rb') as f:
            df = pickle.load(f)
        df['uid'] = df['지역'].astype(str) + '|' + df['자치단체'].astype(str) + '|' \
                    + df['회계'].astype(str) + '|' + df['사업명'].astype(str)
        return df

    # zip 자동 해제
    for z in glob.glob(f'{DATA_DIR}/dataset_*.zip'):
        with zipfile.ZipFile(z) as f:
            f.extractall(DATA_DIR)

    from openpyxl import load_workbook
    tmp_dir = os.path.join(DATA_DIR, '_tmp_years')
    os.makedirs(tmp_dir, exist_ok=True)
    tmp_paths = []
    for path in sorted(glob.glob(f'{DATA_DIR}/*_세부사업별세출현황.xlsx')):
        year = int(os.path.basename(path)[:4])
        print(f'로드 중: {year} ...', flush=True)
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        recs = []
        for r in ws.iter_rows(min_row=3, values_only=True):
            if r[0] is None or r[3] is None:
                continue
            recs.append((r[0], r[1], r[2], r[3], r[4], r[9], r[11], r[12]))
        wb.close()
        d = pd.DataFrame(recs, columns=['지역', '자치단체', '회계', '사업명',
                                        '예산현액', '지출액', '분야', '부문'])
        del recs
        d['예산현액'] = pd.to_numeric(d['예산현액'], errors='coerce')
        d['지출액'] = pd.to_numeric(d['지출액'], errors='coerce')
        d = d.dropna(subset=['예산현액', '지출액'])
        d = d[d['예산현액'] >= 1_000_000]
        # 연도 내에서 먼저 집계해 행 수를 줄인다
        d = d.groupby(['지역', '자치단체', '회계', '사업명'], as_index=False).agg(
            {'예산현액': 'sum', '지출액': 'sum', '분야': 'first', '부문': 'first'})
        d['연도'] = year
        tp = os.path.join(tmp_dir, f'{year}.pkl')
        d.to_pickle(tp)
        tmp_paths.append(tp)
        del d

    df = pd.concat([pd.read_pickle(tp) for tp in tmp_paths], ignore_index=True)
    df['집행률'] = (df['지출액'] / df['예산현액']).clip(0, 1)
    df['uid'] = df['지역'] + '|' + df['자치단체'] + '|' + df['회계'].astype(str) + '|' + df['사업명']
    df.to_pickle(CACHE)
    # 임시 파일 정리
    for tp in tmp_paths:
        os.remove(tp)
    os.rmdir(tmp_dir)
    return df


# ---------- 평가 ----------
def rmse(e):
    e = e.dropna()
    return float(np.sqrt((e ** 2).mean())) if len(e) else np.nan


def main():
    os.makedirs(RESULT_DIR, exist_ok=True)
    df = load_panel()

    # SABER: 검증 연도로 축소계수 추정 후 시험 연도 예측
    saber = SaberModel().fit(df, TRAIN_END, VAL_YEAR)
    test = saber.predict(df, VAL_YEAR, TEST_YEAR)
    # 기준선 예측 결합
    test = predict_baselines(test, df[df.연도 <= VAL_YEAR])

    preds = {'SABER(계층 모형)': 'pred_saber',
             '기존① 개별 추정': 'pred_individual',
             '기존② 전년도': 'pred_last_year',
             '전체 평균': 'pred_global'}
    for name, col in preds.items():
        test[f'err_{col}'] = (test[col] - test['집행률']).abs()

    # 표 1: 이력 구간별 RMSE
    rows = []
    for (a, b), lab in zip(BINS, LABELS):
        m = (test.n >= a) & (test.n <= b)
        rows.append({'이력': lab, 'N': int(m.sum()),
                     **{name: rmse(test.loc[m, f'err_{col}']) for name, col in preds.items()}})
    t1 = pd.DataFrame(rows).round(4)
    t1.to_csv(f'{RESULT_DIR}/table1_rmse_by_history.csv', index=False, encoding='utf-8-sig')
    print('\n[이력 구간별 RMSE]\n', t1.to_string(index=False))

    # 표 2: 종합·대규모(10억+) RMSE
    h = test.n >= 1
    big = test['예산현액'] >= 1e9
    t2 = pd.DataFrame([
        {'구분': '이력 보유 전체', **{n: rmse(test.loc[h, f'err_{c}']) for n, c in preds.items()}},
        {'구분': '대규모(10억+)', **{n: rmse(test.loc[h & big, f'err_{c}']) for n, c in preds.items()}},
    ]).round(4)
    t2.to_csv(f'{RESULT_DIR}/table2_rmse_overall.csv', index=False, encoding='utf-8-sig')
    print('\n[종합 RMSE]\n', t2.to_string(index=False))

    # 표 3: 분야별 RMSE (SABER vs 기존 최선)
    rows = []
    for f in test['분야'].dropna().unique():
        m = (test['분야'] == f) & h
        if m.sum() < 5000:
            continue
        s = rmse(test.loc[m, 'err_pred_saber'])
        base = min(rmse(test.loc[m, 'err_pred_individual']), rmse(test.loc[m, 'err_pred_last_year']))
        rows.append({'분야': f, 'N': int(m.sum()), 'SABER': s, '기존(최선)': base,
                     '개선율%': round((1 - s / base) * 100, 1)})
    t3 = pd.DataFrame(rows).sort_values('개선율%', ascending=False).round(4)
    t3.to_csv(f'{RESULT_DIR}/table3_rmse_by_field.csv', index=False, encoding='utf-8-sig')
    print('\n[분야별 RMSE]\n', t3.to_string(index=False))

    # 신규 사업 커버리지
    n0 = ~h
    print(f"\n[콜드 스타트] 신규 사업 {n0.mean()*100:.1f}% — 기존 방식 예측 불가, "
          f"SABER RMSE {rmse(test.loc[n0,'err_pred_saber']):.4f} "
          f"vs 전체평균 {rmse(test.loc[n0,'err_pred_global']):.4f}")

    # 그림 (matplotlib 설치 시)
    try:
        import matplotlib
        matplotlib.use('Agg')
        matplotlib.rcParams['font.family'] = ['Noto Sans CJK JP', 'Malgun Gothic', 'AppleGothic', 'DejaVu Sans']
        import matplotlib.pyplot as plt
        fig, ax = plt.subplots(figsize=(9, 5))
        x = np.arange(len(BINS)); w = 0.2
        colors = ['#4c72b0', '#c44e52', '#dd8452', '#999999']
        for j, (name, col) in enumerate(preds.items()):
            v = [t1.iloc[i][name] for i in range(len(BINS))]
            ax.bar(x + (j - 1.5) * w, v, w, label=name, color=colors[j])
        ax.set_xticks(x); ax.set_xticklabels(LABELS)
        ax.set_ylabel('RMSE'); ax.legend(); ax.grid(axis='y', alpha=0.3)
        ax.set_title(f'{TEST_YEAR}년 백테스트: 이력 구간별 예측오차')
        fig.tight_layout(); fig.savefig(f'{RESULT_DIR}/fig_rmse_by_history.png', dpi=150)
        print(f'\n그림 저장: {RESULT_DIR}/fig_rmse_by_history.png')
    except ImportError:
        print('matplotlib 미설치 — 그림 생략')


if __name__ == '__main__':
    main()
